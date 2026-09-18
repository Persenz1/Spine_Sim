"""Summarize the completed local boundary study without further contact solves."""
import argparse
import csv
import json
import math
import shutil
import statistics
from collections import Counter
from pathlib import Path

DEFAULT_OUT = Path(r'E:\TestData\IJMS\local_feedback_20260918')
PAPER = Path(r'D:\Pdrive\Document\论文\IJMS')
MATERIALS = {'P240': 'P240', 'P100': 'P100', 'P40': 'P40',
             'fired_brick_standard': '红砖', 'rough_wall': '混凝土'}
EFFECTS = {'YZ': '释放Y和Z：A−B', 'Z': '固定Y释放Z：C−B',
           'Y': '固定Z释放Y：D−B', 'A_minus_C': '恒P下释放Y：A−C',
           'interaction': '交互项：A−C−D+B'}


def read(path):
    with path.open(encoding='utf-8-sig', newline='') as f:
        return list(csv.DictReader(f))


def write(path, rows):
    fields = list(dict.fromkeys(k for row in rows for k in row))
    with path.open('w', encoding='utf-8-sig', newline='') as f:
        writer = csv.DictWriter(f, fields)
        writer.writeheader()
        writer.writerows(rows)


def groups(effects):
    result = []
    for material, case, selection in sorted({(r['material'], r['case_id'], r['selection']) for r in effects}):
        all_rows = [r for r in effects if (r['material'], r['case_id'], r['selection']) == (material, case, selection)]
        for tag in EFFECTS:
            field = 'deltaT_' + tag + '_N'
            rows = [r for r in all_rows if r['A_status'] == 'IN_RANGE' and r[field] != '']
            values = [float(r[field]) for r in rows]
            positive = negative = unresolved = 0
            for r, value in zip(rows, values):
                # An observed replay difference is not a rigorous error bar.
                resolution = max(abs(float(r['replay_T_error_N'])), 16 * math.ulp(max(1., abs(float(r['A_T_N'])))))
                if abs(value) <= resolution:
                    unresolved += 1
                elif value > 0:
                    positive += 1
                else:
                    negative += 1
            result.append(dict(material=material, case_id=case, selection=selection, effect=tag,
                               planned=len(all_rows), valid=len(rows), missing=len(all_rows)-len(rows),
                               median_N=statistics.median(values) if values else None,
                               min_N=min(values) if values else None, max_N=max(values) if values else None,
                               positive=positive, negative=negative, unresolved_vs_A_replay=unresolved))
    return result


def value(text):
    if text == '' or text is None:
        return None
    if isinstance(text, str):
        if text in ('True', 'False'):
            return text == 'True'
        try:
            return float(text) if any(c in text for c in '.eE') else int(text)
        except ValueError:
            return text
    return text


def workbook(out, states, effects, summary):
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    wb = Workbook()
    wb.remove(wb.active)
    notes = [
        ('范围', '50张原表面，每类10张；500基态，四边界共2000目标。'),
        ('A', '自由Y、恒总P；从原接受前态复算原ΔX。'),
        ('B', '固定前态Y和Z，输出夹具反力与实际P。'),
        ('C', '固定前态Y，Z随动且恒P。'),
        ('D', 'Y自由平衡，固定前态Z，输出实际P。'),
        ('失败', '41个目标数值未收敛，保留实际状态；不把失败的差值填0。'),
        ('A特殊例', 'P40_s0009_F3_P1_low原记录采用LSMR recovery；规定recover=False下A未复现。该基态不进入反馈汇总。'),
        ('符号', '正值为相对于该夹持参照增加局部抗拖力，负值为降低；不是全程优劣排名。'),
        ('符号分辨', 'unresolved_vs_A_replay只比较当前ΔT与该基态A总力复算差及浮点尺度，不是严格误差条。'),
        ('固定Z', 'B/D的实际P可以偏离P0甚至为负；夹具额外反力Rz=P0−Pactual，不是恒预载公平比较。'),
        ('重复单位', '表面是重复单位；同表面两种选点及四种边界不是独立样本。'),
        ('原始向量', str(out / 'states') + '，每基态一个JSON，含原前态、记录后态、四种解和逐针分解。'),
        ('地形', str(out / 'surfaces') + '，50份有效包络及原参数；CUDA重建与原raw/envelope标识一致。'),
        ('收敛', '原残差容差10^-3，max_nfev=80，不增加子步、不恢复搜索、不放宽容差。'),
    ]
    sheets = [('说明', ['项目', '内容'], notes)]
    a = [r for r in states if r['variant'] == 'A']
    for title, rows in [('条件汇总', summary), ('500基态', effects), ('2000目标', states), ('A复算', a)]:
        fields = list(dict.fromkeys(k for row in rows for k in row))
        sheets.append((title, fields, [[value(r.get(k)) for k in fields] for r in rows]))
    for title, headers, data in sheets:
        ws = wb.create_sheet(title)
        ws.append(headers)
        for row in data:
            ws.append(row)
        ws.freeze_panes = 'A2'
        ws.auto_filter.ref = ws.dimensions
        for c in ws[1]:
            c.font = Font(bold=True, color='FFFFFF')
            c.fill = PatternFill('solid', fgColor='254E70')
            c.alignment = Alignment(wrap_text=True, vertical='center')
        ws.row_dimensions[1].height = 42
        from openpyxl.utils import get_column_letter
        for j, header in enumerate(headers, 1):
            ws.column_dimensions[get_column_letter(j)].width = min(38, max(15, len(header) + 2))
        if title == '说明':
            ws.column_dimensions['B'].width = 115
            for row in ws.iter_rows(min_row=2):
                row[1].alignment = Alignment(wrap_text=True, vertical='top')
                ws.row_dimensions[row[0].row].height = 34
    path = PAPER / '数据分析和报告/附件/共同背板局部边界试验_20260918.xlsx'
    wb.save(path)
    check = load_workbook(path, read_only=True, data_only=True)
    assert check['2000目标'].max_row == 2001
    assert check['500基态'].max_row == 501
    check.close()
    return path


def update_report(out, states, effects, summary, xlsx):
    a = [r for r in states if r['variant'] == 'A' and r['valid'] == 'True']
    max_replay = max(abs(float(r['replay_T_error_N'])) for r in a)
    counts = {v: sum(r['variant'] == v and r['valid'] == 'True' for r in states) for v in 'ABCD'}
    complete = sum(r['all_valid'] == 'True' for r in effects)
    text = ['## 16. 本机共同背板局部边界试验', '',
            '2026-09-18已按正式输入表完成全部目标。先在RTX 4060 Ti上生成并保存50张原表面的有效包络，原始地形和包络均与返回批次标识完全匹配；正式求解只读取已保存包络，未使用边生成边求解的流水线。', '',
            f'500个基态共2,000个目标均已尝试，1,959个满足原残差和轴向范围条件，41个数值未收敛。A/B/C/D有效数分别为{counts["A"]}/{counts["B"]}/{counts["C"]}/{counts["D"]}，每类目标数均为500；{complete}个基态四边界齐全。求解阶段约62秒。失败保留，不通过换点、子步、恢复搜索或放宽容差补齐。', '',
            f'499个有效A复算均保持记录中的接触／滑动集合，最大总力复算差为{max_replay:.3g} N。唯一未复现项为`P40_s0009_F3_P1_low`：原记录明确使用`recovery_method=lsmr`；规定的`recover=False`普通dense复算未收敛（残差0.01070，原阈值0.001）。该基态不进入反馈汇总，未为追齐记录追加恢复求解。', '',
            '以下是所选低力原步的A−B：恢复自由Y/恒P与固定前态Y/Z之间的局部差值。每组最多10张面，正值为相对于夹持参照增加抗拖力，负值为降低。B的实际法向载荷可改变，不能当作长期恒预载性能排名。', '',
            '| 表面 | 条件 | 有效面数 | ΔT_YZ中位/N | 范围/N | 正/负/本次分辨不清 |',
            '|---|---|---|---|---|---|']
    selected = [('P40', 'F1_P1'), ('P40', 'F2_P1'), ('P100', 'F3_P1'), ('P100', 'F2_P1'),
                ('fired_brick_standard', 'F7_P1'), ('fired_brick_standard', 'F7_P2')]
    for mat, case in selected:
        r = next(r for r in summary if (r['material'], r['case_id'], r['selection'], r['effect']) == (mat, case, 'low', 'YZ'))
        text.append(f"| {MATERIALS[mat]} | {case} | {r['valid']}/10 | {r['median_N']:.6f} | {r['min_N']:.6f}～{r['max_N']:.6f} | {r['positive']}/{r['negative']}/{r['unresolved_vs_A_replay']} |")
    text += ['', '这些低力状态多数表现为随动相对于夹持参照降低当步抗拖力，不能把共同背板天然解释为补谷器。这里的low是各构型自身合格步中的低力位置，并非同X跨构型干预；不能据该表把整个预载救回或增针效应归因给随动。高力状态、C/D与恒P的A−C对照均完整保存在汇总表中，结论限于这些同前史有限步边界干预。', '',
             '符号“本次分辨不清”比较ΔT与该基态A总力复算差及浮点尺度，不是严格数值误差条。B/D实际P最小约−6.557 N、最大约2.043 N，均保留原值，外部Z夹具承担P0−Pactual；负P不自动当作数值失败。', '',
             '逐针法向力幅、法向方向、支撑集合、阈下余项与切向项合计回到ΔT，最大闭合误差约3.11×10⁻¹⁵ N；背板位移＋回缩分解的球心位移闭合误差约2×10⁻¹⁸ m，四边界交互恒等式误差约3.55×10⁻¹⁵ N。上述闭合验证核算，不替代机制解释。', '',
             f'易读表：[共同背板局部边界试验_20260918.xlsx](附件/{xlsx.name})。原始状态与CSV位于`{out}`；`local_states.csv`为2,000个目标，`local_effects.csv`为500个基态，`local_summary.csv`按材料／构型／low-high分别汇总10张面的各项边界差。', '',
             '复现入口：`D:\\Code\\Spine_Sim\\scripts\\run_ijms_local_feedback.py`，`--prepare-only`仅准备全部地形，`--run`只读已落盘地形并执行正式表。已保存的基态JSON按原结果复用，不自动重算失败。运行配置及本次源码副本保存在结果目录。', '']
    path = PAPER / '数据分析和报告/最终确认分析总报告.md'
    report = path.read_text(encoding='utf-8').split('## 16.')[0].rstrip() + '\n\n' + '\n'.join(text)
    report = report.replace('本轮整理没有新增接触求解。', '本轮后处理修正完成，随后完成第16节的2,000个局部接触目标。', 1)
    report = report.replace('最后局部补算已按每类10张原表面准备，共50张、500个基态、2,000个局部目标，尚未执行。', '最后局部补算已在本机完成：50张原表面、500个基态、2,000个局部目标，1,959个有效、41个数值未收敛；结果见第16节。')
    report = report.replace('具体随动因果作用仍待局部试验。', '同前史有限步边界作用已完成，见第16节。')
    report = report.replace('局部试验已改为每类10张，共500基态、2,000个目标，尚未执行。', '局部试验已完成每类10张、500基态、2,000个目标，见第16节。')
    report = report.replace('尚未执行新增接触求解，也没有确定实体实验配置或冻结文章结构。', '已执行第16节局部接触求解，未确定实体实验配置或冻结文章结构。')
    report = report.replace('局部补算已按作者确定的规模准备：', '局部补算已按作者确定的规模执行：')
    report = report.replace('当前求解器尚未因该方案修改，接触补算未执行。', '求解器已增加可选固定Y/Z边界，接触补算已完成，见第16节。')
    report = report.replace('本次仅修正后处理，未执行新的接触求解或共同背板A/B/C/D试验；误差通道诊断不等于随动因果识别。', '本节仅为后处理修正；随后独立完成的共同背板A/B/C/D试验见第16节。误差通道诊断本身不等于随动因果识别。')
    path.write_text(report, encoding='utf-8')


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--output', type=Path, default=DEFAULT_OUT)
    args = parser.parse_args()
    states = read(args.output / 'local_states.csv')
    effects = read(args.output / 'local_effects.csv')
    summary = groups(effects)
    write(args.output / 'local_summary.csv', summary)
    xlsx = workbook(args.output, states, effects, summary)
    update_report(args.output, states, effects, summary, xlsx)
    reproduction = args.output / 'reproduction'
    reproduction.mkdir(exist_ok=True)
    program = Path(__file__).resolve().parents[1]
    for relative in ['scripts/run_ijms_local_feedback.py', 'scripts/summarize_ijms_local_feedback.py', 'src/spine_sim/balanced_contact.py']:
        shutil.copy2(program / relative, reproduction / Path(relative).name)
    print(json.dumps(dict(targets=len(states), status=dict(Counter(r['status'] for r in states)), summary_rows=len(summary), workbook=str(xlsx)), ensure_ascii=False))


if __name__ == '__main__':
    main()
